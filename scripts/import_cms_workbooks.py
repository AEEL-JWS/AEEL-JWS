"""Import AEEL CMS Import worksheets as individual Astro content entries.

Usage: python scripts/import_cms_workbooks.py PATH_TO_WORKBOOK_DIRECTORY
The source workbooks are read only and are intentionally not copied into Git.
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = Path(sys.argv[1])


def rows(filename: str):
    sheet = load_workbook(WORKBOOKS / filename, read_only=True, data_only=True)["CMS Import"]
    values = iter(sheet.values)
    next(values)
    return [row for row in values if any(value is not None for value in row)]


def slug(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def write(collection: str, filename: str, fields: dict):
    destination = ROOT / "src" / "content" / collection / filename
    destination.parent.mkdir(parents=True, exist_ok=True)
    lines = ["---"]
    for key, value in fields.items():
        if value is None or value == "":
            continue
        lines.append(f"{key}: {json.dumps(value, ensure_ascii=False)}")
    lines += ["---", ""]
    destination.write_text("\n".join(lines), encoding="utf-8")


def replace_collection(collection: str):
    directory = ROOT / "src" / "content" / collection
    directory.mkdir(parents=True, exist_ok=True)
    for old in directory.glob("*.md"):
        old.unlink()


publication_book = load_workbook(WORKBOOKS / "AEEL_164_publications_CMS_DOI_master.xlsx", read_only=True, data_only=True)
audit_by_order = {row[0]: row for row in list(publication_book["Master Audit"].values)[1:]}
cover_by_order = {order: row[10].title() for order, row in audit_by_order.items() if isinstance(row[10], str) and "cover" in row[10].lower()}
papers = rows("AEEL_164_publications_CMS_DOI_master.xlsx")
assert len(papers) == 164 and len({r[0].casefold() for r in papers}) == 164
assert len(audit_by_order) == len(papers)
replace_collection("publications")
for index, row in enumerate(papers, 1):
    title, authors, journal, year, doi, link, area, featured, image = row
    audit = audit_by_order[index]
    assert audit[4] == title and audit[1] == year
    bibliography = str(audit[3]).strip().rstrip(".") if audit[3] is not None else None
    if bibliography:
        bibliography = re.sub(r"(?<=[A-Za-z0-9])\s*-\s*(?=[A-Za-z0-9])", "–", bibliography)
        bibliography = re.sub(r"^(\d+)\((\d+)\)\s+", r"\1, \2, ", bibliography)
        bibliography = re.sub(r"^(\d+)[:.]\s*(\d+)", r"\1, \2", bibliography)
        bibliography = re.sub(r"^(\d+)\s+(\d+)", r"\1, \2", bibliography)
        bibliography = re.sub(r"(,\s*\d+)\s+(e\d+)$", r"\1, \2", bibliography)
    status = audit[10] if audit[10] in {"Accepted", "In press", "Early View", "ASAP"} else None
    write("publications", f"{year}-{index:03d}-{slug(title)[:65]}.md", {
        "title": title, "authors": authors, "journal": journal, "year": year,
        "bibliography": bibliography, "status": status,
        "doi": doi, "link": link, "area": area, "featured": bool(featured),
        "image": image, "coverType": cover_by_order.get(index), "order": index,
    })

members = rows("AEEL_members_CMS_master.xlsx")
assert len(members) == 17 and len({r[1].casefold() for r in members}) == 17
replace_collection("people")
for row in members:
    order, name, role, affiliation, email, interests, image_url, image_name, _, active, notes = row
    write("people", f"{slug(name)}.md", {
        "name": name, "role": role, "affiliation": affiliation, "email": email,
        "interests": interests, "image": f"/images/members/{Path(image_name).stem}.webp" if image_url else None,
        "sourceImageUrl": image_url, "order": order, "active": bool(active),
    })

alumni = rows("AEEL_alumni_CMS_master.xlsx")
assert len(alumni) == 14 and len({r[1].casefold() for r in alumni}) == 14
replace_collection("alumni")
for row in alumni:
    order, name, category, source_category, year, month, status, position, organization, current_text, active, notes = row
    write("alumni", f"{slug(name)}.md", {
        "name": name, "category": category, "graduationYear": year,
        "graduationMonth": month, "currentStatus": status,
        "currentPosition": position, "currentOrganization": organization,
        "currentText": current_text, "order": order, "active": bool(active),
    })

notices = rows("AEEL_notices_134_CMS_master.xlsx")
assert len(notices) == 134 and len({r[7] for r in notices}) == 134
replace_collection("notices")
for row in notices:
    order, date, year, category, source_category, title, source_title, source_url, source_slug, views, featured, active, image, body, notes = row
    write("notices", f"{date:%Y-%m-%d}-{order:03d}.md", {
        "title": title, "date": date.date().isoformat(), "category": category,
        "sourceTitle": source_title, "sourceUrl": source_url,
        "legacyViews": views, "featured": bool(featured), "active": bool(active),
        "image": image, "body": body,
    })

print(f"Imported {len(papers)} publications, {len(members)} members, {len(alumni)} alumni, {len(notices)} notices.")
